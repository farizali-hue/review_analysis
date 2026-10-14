from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(r"C:\Users\dell\OneDrive - Utopia\hassan\Creatine-HCL\claims\claims-creatine_hcl.xlsx")
OUT = Path(r"D:\analysis\creatine_hcl_rows.json")


def clean_text(value):
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""
    if any(marker in text for marker in ("â", "Ã", "�")):
        try:
            text = text.encode("cp1252", errors="ignore").decode("utf-8", errors="ignore")
        except Exception:
            pass
    return (
        text.replace("\u2013", "-")
        .replace("\u2014", "-")
        .replace("\u2018", "'")
        .replace("\u2019", "'")
        .replace("\u201c", '"')
        .replace("\u201d", '"')
        .replace("\xa0", " ")
    )


def main() -> None:
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [clean_text(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]

    rows = []
    for row_idx in range(2, ws.max_row + 1):
        record = {}
        for col_idx, header in enumerate(headers, 1):
            record[header] = clean_text(ws.cell(row_idx, col_idx).value)
        if record.get("asin") or record.get("url"):
            record["_source_row"] = row_idx
            rows.append(record)

    OUT.write_text(json.dumps(rows, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({"rows": len(rows), "output": str(OUT)}, indent=2))


if __name__ == "__main__":
    main()
