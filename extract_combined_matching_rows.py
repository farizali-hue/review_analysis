import csv
from pathlib import Path

import openpyxl


INPUT_XLSX = Path(r"C:\Users\dell\OneDrive - Utopia\maha\combined_matching_files.xlsx")
OUT_CSV = Path(r"D:\analysis\combined_matching_rows.csv")


def main() -> None:
    wb = openpyxl.load_workbook(INPUT_XLSX, data_only=True)
    ws = wb["combined_claims"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    headers = [h if h else f"blank_{i}" for i, h in enumerate(headers, 1)]

    rows = []
    for r in range(2, ws.max_row + 1):
        row = dict(zip(headers, [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]))
        if row.get("asin") and row.get("url"):
            rows.append(row)

    with OUT_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)

    print(f"wrote {len(rows)} rows to {OUT_CSV}")


if __name__ == "__main__":
    main()
